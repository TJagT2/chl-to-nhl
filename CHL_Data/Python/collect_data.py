"""
Download all the outside data the project uses. Run this first.

    1. CHL standings 1985-86 to 2015-16 (Wikipedia season pages): every team's
       record, goals for/against, division and conference.
       -> CHL_Data_Full/Standings/CHL Standings 1985-2016.xlsx
    2. NHL careers 1986-87 to 2025-26 (NHL stats API): every skater's games, goals,
       assists, points, birth date, height and draft position.
       -> CHL_Data_Full/NHL/NHL Careers 1986-2026.xlsx
    3. NHL draft picks 1986-2016 (NHL records API): every pick, including players
       who never played in the NHL, with amateur team, height and birth date.
       -> CHL_Data_Full/NHL/NHL Draft Picks 1986-2016.xlsx

Run from VS Code (open the file, click Run) or a terminal:
    python collect_data.py                 # all three
    python collect_data.py --only draft    # just one: standings, nhl or draft

Downloads are cached in CHL_Data_Full/_cache, so re-running is quick and doesn't
contact the sites again. Delete a cache folder to force a fresh download.
Wikipedia content is CC BY-SA: credit it as a source in your write-up.
"""

import argparse
import json
import re
import sys
import time
from collections import Counter, defaultdict

import lxml.html
import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from common import CACHE_DIR, DRAFT_FILE, LEAGUES, NHL_CAREERS_FILE, STANDINGS_FILE

FIRST_SEASON, LAST_SEASON = 1985, 2015          # CHL seasons 1985-86 ... 2015-16
FIRST_DRAFT, LAST_DRAFT = 1986, 2016
NHL_SEASONS = (19861987, 20252026)

# Wikipedia asks scripts to identify themselves with a way to contact the owner.
# Optional: put your email or a GitHub URL here, e.g. CONTACT = "you@example.com"
CONTACT = ""
USER_AGENT = ("CHL-NHL-research/2.0 (personal hockey research project"
              + (f"; {CONTACT}" if CONTACT else "") + ")")


def write_excel(path, header, rows, sheet_title, widths=None, extra_sheets=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheets = [(sheet_title, header, rows, widths)] + list(extra_sheets or [])
    for i, (title, hdr, body, w) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(hdr)
        for r in body:
            # strip characters Excel can't store (some birthplaces contain them)
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in r])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for j, name in enumerate(hdr, start=1):
            ws.column_dimensions[get_column_letter(j)].width = (w or {}).get(name, 12)
    wb.save(path)


# =====================================================================
# 1. CHL standings (Wikipedia)
# =====================================================================

WIKI_API = "https://en.wikipedia.org/w/api.php"
WIKI_DELAY_SECONDS = 3.0     # pause between downloads
MAX_RETRIES = 5              # if Wikipedia says "too many requests", wait and try again
WIKI_CACHE = CACHE_DIR / "wikipedia"

# Column headings used on the pages, mapped to one standard name
HEADER_NAMES = {
    "PLD": "GP", "GP": "GP", "W": "W", "L": "L", "T": "T",
    "OTL": "OTL", "OL": "OTL", "SOL": "SOL", "SL": "SOL",
    "PTS": "PTS", "GF": "GF", "GA": "GA",
    "TEAM": "TEAM", "DIV": "DIV", "POS": "POS", "RANK": "RANK",
}
STATS = ["GP", "W", "L", "T", "OTL", "SOL", "PTS", "GF", "GA"]
DIVISION_ABBREVIATIONS = {"CEN": "Central", "CEN.": "Central"}

# Seasons where Wikipedia shows divisions but not which conference they belong
# to (WHL wild-card layout). WHL conferences: East + Central = Eastern, B.C. + U.S. = Western.
WHL_CONFERENCES = {"East": "Eastern Conference", "Central": "Eastern Conference",
                   "B.C.": "Western Conference", "U.S.": "Western Conference"}

COLUMNS = ["season", "league", "team", "conference", "division", "gp", "w", "l", "t",
           "otl", "sol", "pts", "gf", "ga", "gf_per_game", "source"]


def page_title(league, start_year):
    end = "2000" if start_year == 1999 else str(start_year + 1)[-2:]
    return f"{start_year}–{end} {league} season"          # en dash, as Wikipedia uses


def fetch_page_html(session, title):
    """Return the page's HTML via the Wikipedia API, using the local cache if possible."""
    cache_file = WIKI_CACHE / (re.sub(r"[^\w-]+", "_", title) + ".html")
    if cache_file.exists():
        return cache_file.read_text(encoding="utf-8"), False
    params = {"action": "parse", "page": title, "prop": "text", "redirects": 1,
              "format": "json", "formatversion": 2}
    for attempt in range(1, MAX_RETRIES + 1):
        response = session.get(WIKI_API, params=params, timeout=30)
        if response.status_code != 429:
            break
        # Rate limited: wait as long as Wikipedia asks (or 30s, 60s, 120s ...)
        retry_after = response.headers.get("Retry-After", "")
        wait = int(retry_after) if retry_after.isdigit() else 30 * 2 ** (attempt - 1)
        print(f"   Wikipedia asked us to slow down; waiting {wait}s (attempt {attempt}/{MAX_RETRIES})")
        time.sleep(wait)
    response.raise_for_status()
    data = response.json()
    if "error" in data:
        raise RuntimeError(f"Wikipedia: {data['error'].get('info', data['error'])}")
    html = data["parse"]["text"]
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(html, encoding="utf-8")
    return html, True


def clean(text):
    return re.sub(r"\s+", " ", (text or "").replace("\xa0", " ")).strip()


def clean_team(name):
    """Remove playoff markers like 'x – ', 'y-', 'x, z – ', 'y ' and a trailing ' z'."""
    name = clean(name)
    name = re.sub(r"^[a-zA-Z](?:\s*,\s*[a-zA-Z])*\s*[–—-]\s*(?=[A-Z0-9])", "", name)
    name = re.sub(r"^[a-z]\s+(?=[A-Z])", "", name)
    name = re.sub(r"\s+[a-zA-Z]$", "", name)
    name = name.replace("–", "-")             # "Acadie–Bathurst" -> "Acadie-Bathurst"
    return name.strip(" *†‡")


def to_int(text):
    text = clean(text).replace("−", "-")
    return int(text) if re.fullmatch(r"[+-]?\d+", text) else None


def division_name(text):
    """'TELUS East Division' -> 'East Division', 'Central Division: Top 3' -> 'Central Division'."""
    text = clean(text).split(":")[0]
    text = DIVISION_ABBREVIATIONS.get(text.upper(), text)
    core = clean(re.sub(r"\bDivision\b", "", text, flags=re.I))
    without_sponsor = clean(re.sub(r"\bTELUS\b", "", core, flags=re.I))
    core = without_sponsor or core              # "TELUS East" -> "East", but keep "Telus" alone
    return f"{core} Division" if core else None


def label_kind(text):
    """Classify a heading/caption/label as 'conference', 'division' or None."""
    t = clean(text).lower()
    if not t or "standings" in t or "wild card" in t or t in ("overall", "team"):
        return None
    if "conference" in t:
        return "conference"
    if "division" in t:
        return "division"
    return None


def parse_standings(html):
    """
    Return a list of team dicts from one season page.

    Walks the page top to bottom, remembering the most recent headings, and
    reads every table that has GF and GA columns. Handles the different layouts
    used over the years: division name as a caption, as a heading, as the first
    column heading, or in a DIV column; header rows repeated inside one table;
    and the same team appearing in both a division and a conference table.
    """
    doc = lxml.html.fromstring(html)
    for el in doc.xpath("//sup | //style | //span[contains(@class,'mw-editsection')]"):
        el.drop_tree()                     # footnote markers, CSS, '[edit]' links

    headings = {}                          # level -> text
    rows = []
    for el in doc.iter("h2", "h3", "h4", "h5", "table"):
        if el.tag != "table":
            level = int(el.tag[1])
            headings = {k: v for k, v in headings.items() if k < level}
            headings[level] = clean(el.text_content())
            continue
        if "playoff" in headings.get(2, "").lower():
            continue                       # playoff round-robin tables, not the regular season
        if el.xpath(".//table"):
            continue                       # a layout table wrapping other tables
        rows.extend(parse_table(el, headings))
    return rows


def parse_table(table, headings):
    caption = clean(" ".join(table.xpath("./caption//text()")))
    header, row_label, out = None, None, []

    for tr in table.xpath("./tr | ./thead/tr | ./tbody/tr"):
        texts = [clean(c.text_content()) for c in tr.xpath("./th | ./td")]
        upper = [t.upper() for t in texts]
        if "GF" in upper and "GA" in upper:                     # a header row
            header = [HEADER_NAMES.get(t, t) for t in upper]
            # Some tables have no "Team" heading: the team column is headed by
            # the division name instead (e.g. "East Division"), sometimes after a Rank column.
            label_cols = [i for i, t in enumerate(upper) if t not in HEADER_NAMES]
            row_label = texts[label_cols[0]] if label_cols else None
            continue
        if header is None or not any(texts):
            continue

        # Line cells up with the header from the right (stats are always at the end)
        n = min(len(texts), len(header))
        cells = dict(zip(header[-n:], texts[-n:]))
        team_col = "TEAM" if "TEAM" in header else (row_label or "").upper()
        team = clean_team(cells.get(team_col, ""))
        gp = to_int(cells.get("GP", ""))
        if not team or gp is None or to_int(team) is not None:
            continue

        labels = [cells.get("DIV"), row_label, caption, headings.get(5),
                  headings.get(4), headings.get(3)]
        division = conference = None
        if cells.get("DIV"):
            division = division_name(cells["DIV"])
        for label in labels[1:]:
            kind = label_kind(label or "")
            if kind == "division" and not division:
                division = division_name(label)
            elif kind == "conference" and not conference:
                conference = clean(label).split(":")[0].title()
        out.append({"team": team, "division": division, "conference": conference,
                    **{k.lower(): to_int(cells.get(k, "")) for k in STATS}})
    return out


def combine(rows, league, start_year):
    """Merge repeat listings of a team, drop non-regular-season rows, fill conferences."""
    if not rows:
        return [], []
    warnings = []
    max_gp = max(r["gp"] for r in rows)
    teams = {}
    for r in rows:
        if r["gp"] < 0.6 * max_gp:          # e.g. a 4-game round robin
            continue
        key = r["team"].casefold()
        if key not in teams:
            teams[key] = dict(r)
            continue
        t = teams[key]
        if r["gf"] is not None and t["gf"] is not None and r["gf"] != t["gf"]:
            warnings.append(f"{r['team']}: goals for listed as {t['gf']} and {r['gf']}")
        for field in ("division", "conference"):
            t[field] = t[field] or r[field]
        for k in STATS:
            t[k.lower()] = t[k.lower()] if t[k.lower()] is not None else r[k.lower()]

    # Fill a missing conference from other teams in the same division that season
    div_conf = defaultdict(Counter)
    for t in teams.values():
        if t["division"] and t["conference"]:
            div_conf[t["division"]][t["conference"]] += 1
    for t in teams.values():
        if not t["conference"] and t["division"]:
            if div_conf[t["division"]]:
                t["conference"] = div_conf[t["division"]].most_common(1)[0][0]
            elif league == "WHL" and start_year >= 2001:
                t["conference"] = WHL_CONFERENCES.get(t["division"].replace(" Division", ""))
    return list(teams.values()), warnings



def collect_standings(session):
    standings, checks, downloaded = [], [], 0
    for league in LEAGUES:
        for year in range(FIRST_SEASON, LAST_SEASON + 1):
            season = f"{year}-{year + 1}"
            title = page_title(league, year)
            url = "https://en.wikipedia.org/wiki/" + title.replace(" ", "_")
            try:
                html, fresh = fetch_page_html(session, title)
            except Exception as exc:
                print(f"  {league} {season}: could not download ({exc})")
                checks.append([season, league, 0, None, None, None, f"download failed: {exc}"])
                continue
            if fresh:
                downloaded += 1
                time.sleep(WIKI_DELAY_SECONDS)

            teams, warnings = combine(parse_standings(html), league, year)
            for t in sorted(teams, key=lambda t: (t["conference"] or "", t["division"] or "",
                                                  -(t["pts"] or 0))):
                gp, gf = t["gp"], t["gf"]
                standings.append([season, league, t["team"], t["conference"], t["division"]]
                                 + [t[k.lower()] for k in STATS]
                                 + [round(gf / gp, 2) if gp and gf is not None else None, url])
            gf_total = sum(t["gf"] or 0 for t in teams)
            ga_total = sum(t["ga"] or 0 for t in teams)
            if not teams:
                warnings.append("no standings table found")
            if any(t["gf"] is None for t in teams):
                warnings.append("some teams have no goals for")
            if gf_total != ga_total:
                warnings.append(f"league GF ({gf_total}) != GA ({ga_total})")
            if not any(t["division"] or t["conference"] for t in teams):
                warnings.append("no division/conference found")
            checks.append([season, league, len(teams), gf_total, ga_total,
                           gf_total - ga_total, "; ".join(warnings)])
            if warnings:
                print(f"  {league} {season}: {len(teams)} teams  <-- check: {'; '.join(warnings)}")

    write_excel(STANDINGS_FILE, COLUMNS, standings, "Standings",
                {"team": 30, "conference": 20, "division": 20, "source": 50},
                extra_sheets=[("Checks", ["season", "league", "teams", "league_gf", "league_ga",
                                          "gf_minus_ga", "warnings"], checks, {"warnings": 60})])
    flagged = sum(1 for c in checks if c[-1])
    print(f"Standings: {len(standings)} team seasons ({downloaded} pages downloaded); "
          f"{flagged} league seasons noted on the Checks sheet -> {STANDINGS_FILE.name}")


# =====================================================================
# 2. NHL careers (NHL stats API)
# =====================================================================

NHL_BIOS_API = "https://api.nhle.com/stats/rest/en/skater/bios"
NHL_PAGE_SIZE = 100
NHL_CACHE = CACHE_DIR / "nhl"
NHL_FIELDS = [
    ("nhl_player_id", "playerId"), ("full_name", "skaterFullName"), ("last_name", "lastName"),
    ("birth_date", "birthDate"), ("birth_city", "birthCity"),
    ("birth_state_province", "birthStateProvinceCode"), ("birth_country", "birthCountryCode"),
    ("nationality", "nationalityCode"), ("position", "positionCode"), ("shoots", "shootsCatches"),
    ("height_in", "height"), ("weight_lb", "weight"),
    ("draft_year", "draftYear"), ("draft_round", "draftRound"), ("draft_overall", "draftOverall"),
    ("first_nhl_season", "firstSeasonForGameType"),
    ("nhl_gp", "gamesPlayed"), ("nhl_goals", "goals"), ("nhl_assists", "assists"),
    ("nhl_points", "points"), ("hall_of_fame", "isInHallOfFameYn"),
]


def fetch_nhl_page(session, start):
    cache_file = NHL_CACHE / f"bios_{start:05d}.json"
    if cache_file.exists():
        return json.loads(cache_file.read_text(encoding="utf-8")), False
    params = {
        "isAggregate": "true",
        "isGame": "false",
        "sort": json.dumps([{"property": "playerId", "direction": "ASC"}]),   # stable paging
        "start": start,
        "limit": NHL_PAGE_SIZE,
        "cayenneExp": f"gameTypeId=2 and seasonId>={NHL_SEASONS[0]} and seasonId<={NHL_SEASONS[1]}",
    }
    response = session.get(NHL_BIOS_API, params=params, timeout=30)
    if response.status_code in (403, 429):
        sys.exit(f"STOPPED: the NHL API refused the request (HTTP {response.status_code}). "
                 "Wait a while and run again; downloaded pages are cached.")
    response.raise_for_status()
    data = response.json()
    if "data" not in data or "total" not in data:
        sys.exit("STOPPED: the NHL API returned an unexpected format; it may have changed.")
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(json.dumps(data), encoding="utf-8")
    return data, True


def collect_nhl_careers(session):
    players, start, total, downloaded = {}, 0, None, 0
    while total is None or start < total:
        data, fresh = fetch_nhl_page(session, start)
        total = data["total"]
        for p in data["data"]:
            players[p["playerId"]] = p             # keyed by ID, so overlaps can't duplicate
        if fresh:
            downloaded += 1
            time.sleep(1.0)
        start += NHL_PAGE_SIZE
    rows = [[p.get(field) for _, field in NHL_FIELDS]
            for p in sorted(players.values(), key=lambda p: (p.get("lastName") or "", p["playerId"]))]
    write_excel(NHL_CAREERS_FILE, [n for n, _ in NHL_FIELDS], rows, "NHL Skaters",
                {"full_name": 24, "birth_city": 24})
    print(f"NHL careers: {len(players):,} skaters ({downloaded} pages downloaded) -> {NHL_CAREERS_FILE.name}")


# =====================================================================
# 3. NHL draft picks (NHL records API)
# =====================================================================

DRAFT_API = "https://records.nhl.com/site/api/draft"
DRAFT_CACHE = CACHE_DIR / "nhl" / f"draft_{FIRST_DRAFT}_{LAST_DRAFT}.json"
DRAFT_FIELDS = [
    ("draft_year", "draftYear"), ("round", "roundNumber"), ("pick_in_round", "pickInRound"),
    ("overall_pick", "overallPickNumber"), ("drafted_by", "triCode"), ("nhl_player_id", "playerId"),
    ("player", "playerName"), ("position", "position"), ("shoots", "shootsCatches"),
    ("height_in", "height"), ("weight_lb", "weight"), ("birth_date", "birthDate"),
    ("birth_place", "birthPlace"), ("country", "countryCode"),
    ("amateur_team", "amateurClubName"), ("amateur_league", "amateurLeague"),
]


def collect_draft(session):
    if DRAFT_CACHE.exists():
        picks, downloaded = json.loads(DRAFT_CACHE.read_text(encoding="utf-8")), False
    else:
        params = {"cayenneExp": f"draftYear>={FIRST_DRAFT} and draftYear<={LAST_DRAFT}", "limit": -1}
        for _ in range(3):
            response = session.get(DRAFT_API, params=params, timeout=60)
            if response.status_code not in (429, 503):
                break
            time.sleep(30)
        if response.status_code in (403, 429):
            sys.exit(f"STOPPED: the NHL records API refused the request (HTTP {response.status_code}).")
        response.raise_for_status()
        data = response.json()
        if "data" not in data:
            sys.exit("STOPPED: the NHL records API returned an unexpected format.")
        picks, downloaded = data["data"], True
        DRAFT_CACHE.parent.mkdir(parents=True, exist_ok=True)
        DRAFT_CACHE.write_text(json.dumps(picks), encoding="utf-8")
    rows = [[p.get(field) for _, field in DRAFT_FIELDS]
            for p in sorted(picks, key=lambda p: (p["draftYear"], p["overallPickNumber"]))]
    write_excel(DRAFT_FILE, [n for n, _ in DRAFT_FIELDS], rows, "Draft Picks",
                {"player": 24, "amateur_team": 28, "birth_place": 24})
    print(f"Draft: {len(picks):,} picks{' (downloaded)' if downloaded else ''} -> {DRAFT_FILE.name}")


# =====================================================================

def main():
    parser = argparse.ArgumentParser(description="Download standings, NHL careers and draft picks.")
    parser.add_argument("--only", choices=["standings", "nhl", "draft"],
                        help="download just one dataset (default: all three)")
    args = parser.parse_args()

    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    steps = {"standings": collect_standings, "nhl": collect_nhl_careers, "draft": collect_draft}
    for name, step in steps.items():
        if args.only in (None, name):
            step(session)
    print("\nDone. Next: python load_database.py")


if __name__ == "__main__":
    main()
